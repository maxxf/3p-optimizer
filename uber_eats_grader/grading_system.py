import os
import sys
import logging
import json
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from datetime import datetime
from fpdf import FPDF
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    logging.warning("python-pptx not available, presentation generation will be disabled")

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("GradingSystem")

class GradingSystem:
    """
    A class for grading restaurant delivery storefronts based on analysis results.
    """
    
    def __init__(self, output_dir=None):
        """
        Initialize the Grading System.
        
        Args:
            output_dir (str, optional): Directory to save output files
        """
        logger.info("Initializing Grading System...")
        
        # Set output directory
        self.output_dir = output_dir or os.path.join(os.getcwd(), 'output')
        os.makedirs(self.output_dir, exist_ok=True)
        
        # Define grading scales
        self.grade_scale = {
            0.9: 'A+',
            0.85: 'A',
            0.8: 'A-',
            0.75: 'B+',
            0.7: 'B',
            0.65: 'B-',
            0.6: 'C+',
            0.55: 'C',
            0.5: 'C-',
            0.45: 'D+',
            0.4: 'D',
            0.0: 'F'
        }
        
        # Define component weights
        self.component_weights = {
            'menu_structure': 0.25,
            'pricing': 0.20,
            'customer_sentiment': 0.30,
            'visual_presentation': 0.15,
            'promotion_strategy': 0.10
        }
        
        # Load benchmarks
        self.benchmarks = self._load_benchmarks()
        
        # Set color scheme for visualizations
        self.colors = {
            'primary': '#FF5722',  # Spice orange
            'secondary': '#212B38',  # Dark blue
            'accent': '#4CAF50',  # Green
            'light': '#F5F5F5',  # Light gray
            'dark': '#333333',  # Dark gray
            'grades': {
                'A+': '#4CAF50',  # Green
                'A': '#8BC34A',
                'A-': '#CDDC39',
                'B+': '#FFEB3B',  # Yellow
                'B': '#FFC107',
                'B-': '#FF9800',
                'C+': '#FF5722',  # Orange
                'C': '#F44336',
                'C-': '#E91E63',
                'D+': '#9C27B0',  # Purple
                'D': '#673AB7',
                'F': '#F44336'   # Red
            }
        }
    
    def _load_benchmarks(self):
        """
        Load benchmark data for comparison.
        
        Returns:
            dict: Benchmark data
        """
        try:
            # In a real implementation, this would load from a database or file
            # For now, we'll use hardcoded benchmark values
            return {
                'menu_structure': {
                    'avg_score': 0.75,
                    'percentiles': {
                        90: 0.90,
                        75: 0.82,
                        50: 0.75,
                        25: 0.65,
                        10: 0.55
                    }
                },
                'pricing': {
                    'avg_score': 0.70,
                    'percentiles': {
                        90: 0.85,
                        75: 0.78,
                        50: 0.70,
                        25: 0.62,
                        10: 0.55
                    }
                },
                'customer_sentiment': {
                    'avg_score': 0.80,
                    'percentiles': {
                        90: 0.92,
                        75: 0.85,
                        50: 0.80,
                        25: 0.70,
                        10: 0.60
                    }
                },
                'visual_presentation': {
                    'avg_score': 0.65,
                    'percentiles': {
                        90: 0.85,
                        75: 0.75,
                        50: 0.65,
                        25: 0.55,
                        10: 0.45
                    }
                },
                'promotion_strategy': {
                    'avg_score': 0.60,
                    'percentiles': {
                        90: 0.80,
                        75: 0.70,
                        50: 0.60,
                        25: 0.50,
                        10: 0.40
                    }
                }
            }
        except Exception as e:
            logger.error(f"Error loading benchmarks: {e}")
            return {}
    
    def process_analysis_results(self, analysis_results):
        """
        Process analysis results and generate grading.
        
        Args:
            analysis_results (dict): Analysis results from AnalysisEngine
            
        Returns:
            dict: Grading results
        """
        try:
            restaurant_name = analysis_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            logger.info(f"Processing analysis results for {restaurant_name}")
            
            # Extract component scores
            component_scores = {
                'menu_structure': analysis_results.get('menu_structure', {}).get('score', 0),
                'pricing': analysis_results.get('pricing', {}).get('score', 0),
                'customer_sentiment': analysis_results.get('customer_sentiment', {}).get('score', 0),
                'visual_presentation': analysis_results.get('visual_presentation', {}).get('score', 0),
                'promotion_strategy': analysis_results.get('promotion_strategy', {}).get('score', 0)
            }
            
            # Calculate weighted overall score
            weighted_scores = {}
            for component, score in component_scores.items():
                weight = self.component_weights.get(component, 0)
                weighted_scores[component] = score * weight
            
            overall_score = sum(weighted_scores.values())
            
            # Calculate grades
            overall_grade = self.calculate_grade(overall_score)
            component_grades = {component: self.calculate_grade(score) for component, score in component_scores.items()}
            
            # Compare to benchmarks
            benchmark_comparisons = self.compare_to_benchmarks(component_scores)
            
            # Prioritize recommendations
            all_recommendations = []
            for component, component_data in analysis_results.items():
                if component in ['restaurant_info', 'overall_score']:
                    continue
                
                if isinstance(component_data, dict) and 'recommendations' in component_data:
                    for rec in component_data['recommendations']:
                        all_recommendations.append({
                            'component': component,
                            'recommendation': rec,
                            'score': component_scores.get(component, 0)
                        })
            
            prioritized_recommendations = self.prioritize_recommendations(all_recommendations)
            
            # Combine all strengths and weaknesses
            all_strengths = []
            all_weaknesses = []
            
            for component, component_data in analysis_results.items():
                if component in ['restaurant_info', 'overall_score']:
                    continue
                
                if isinstance(component_data, dict):
                    if 'strengths' in component_data:
                        for strength in component_data['strengths']:
                            all_strengths.append({
                                'component': component,
                                'text': strength
                            })
                    
                    if 'weaknesses' in component_data:
                        for weakness in component_data['weaknesses']:
                            all_weaknesses.append({
                                'component': component,
                                'text': weakness
                            })
            
            # Compile grading results
            grading_results = {
                'restaurant_info': analysis_results.get('restaurant_info', {}),
                'overall_score': overall_score,
                'overall_grade': overall_grade,
                'component_scores': component_scores,
                'component_grades': component_grades,
                'benchmark_comparisons': benchmark_comparisons,
                'strengths': all_strengths,
                'weaknesses': all_weaknesses,
                'prioritized_recommendations': prioritized_recommendations,
                'analysis_results': analysis_results
            }
            
            return grading_results
        except Exception as e:
            logger.error(f"Error processing analysis results: {e}")
            return {
                'error': str(e),
                'restaurant_info': analysis_results.get('restaurant_info', {}),
                'overall_grade': 'F',
                'component_grades': {},
                'strengths': [],
                'weaknesses': [],
                'prioritized_recommendations': []
            }
    
    def calculate_grade(self, score):
        """
        Calculate letter grade based on score.
        
        Args:
            score (float): Numerical score (0-1)
            
        Returns:
            str: Letter grade
        """
        for threshold, grade in sorted(self.grade_scale.items(), reverse=True):
            if score >= threshold:
                return grade
        return 'F'
    
    def compare_to_benchmarks(self, component_scores):
        """
        Compare component scores to industry benchmarks.
        
        Args:
            component_scores (dict): Component scores
            
        Returns:
            dict: Benchmark comparison results
        """
        comparisons = {}
        
        for component, score in component_scores.items():
            benchmark = self.benchmarks.get(component, {})
            avg_score = benchmark.get('avg_score', 0.5)
            percentiles = benchmark.get('percentiles', {})
            
            # Determine percentile
            percentile = None
            for p, threshold in sorted(percentiles.items()):
                if score >= threshold:
                    percentile = p
                    break
            
            if percentile is None:
                percentile = 0
            
            # Calculate difference from average
            diff_from_avg = score - avg_score
            
            # Determine position
            if diff_from_avg >= 0.1:
                position = 'above_average'
            elif diff_from_avg <= -0.1:
                position = 'below_average'
            else:
                position = 'average'
            
            comparisons[component] = {
                'score': score,
                'benchmark_avg': avg_score,
                'diff_from_avg': diff_from_avg,
                'percentile': percentile,
                'position': position
            }
        
        return comparisons
    
    def prioritize_recommendations(self, recommendations):
        """
        Prioritize recommendations based on component scores and importance.
        
        Args:
            recommendations (list): All recommendations
            
        Returns:
            list: Prioritized recommendations
        """
        # Sort recommendations by component score (ascending) and then by component weight (descending)
        sorted_recommendations = sorted(
            recommendations,
            key=lambda x: (x['score'], -self.component_weights.get(x['component'], 0))
        )
        
        # Return top 5 recommendations
        return sorted_recommendations[:5]
    
    def generate_visualizations(self, grading_results, output_dir=None):
        """
        Generate visualizations for grading results.
        
        Args:
            grading_results (dict): Grading results
            output_dir (str, optional): Directory to save visualizations
            
        Returns:
            dict: Paths to generated visualization files
        """
        try:
            output_dir = output_dir or self.output_dir
            os.makedirs(output_dir, exist_ok=True)
            
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown')
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            
            # Set style
            plt.style.use('seaborn-v0_8-whitegrid')
            
            # Generate component scores radar chart
            radar_path = os.path.join(output_dir, f"{restaurant_name}_radar_{timestamp}.png")
            self._generate_radar_chart(grading_results, radar_path)
            
            # Generate benchmark comparison bar chart
            benchmark_path = os.path.join(output_dir, f"{restaurant_name}_benchmark_{timestamp}.png")
            self._generate_benchmark_chart(grading_results, benchmark_path)
            
            # Generate grade distribution pie chart
            grade_path = os.path.join(output_dir, f"{restaurant_name}_grades_{timestamp}.png")
            self._generate_grade_chart(grading_results, grade_path)
            
            return {
                'radar_chart': radar_path,
                'benchmark_chart': benchmark_path,
                'grade_chart': grade_path
            }
        except Exception as e:
            logger.error(f"Error generating visualizations: {e}")
            return {}
    
    def _generate_radar_chart(self, grading_results, output_path):
        """
        Generate radar chart for component scores.
        
        Args:
            grading_results (dict): Grading results
            output_path (str): Path to save the chart
        """
        try:
            component_scores = grading_results.get('component_scores', {})
            
            # Prepare data
            categories = list(component_scores.keys())
            categories = [c.replace('_', ' ').title() for c in categories]
            
            values = list(component_scores.values())
            
            # Close the loop
            categories.append(categories[0])
            values.append(values[0])
            
            # Create figure
            fig, ax = plt.subplots(figsize=(10, 8), subplot_kw=dict(polar=True))
            
            # Set background color
            fig.patch.set_facecolor(self.colors['light'])
            ax.set_facecolor(self.colors['light'])
            
            # Plot data
            angles = np.linspace(0, 2*np.pi, len(categories), endpoint=False).tolist()
            angles.append(angles[0])
            
            ax.plot(angles, values, 'o-', linewidth=2, color=self.colors['primary'])
            ax.fill(angles, values, alpha=0.25, color=self.colors['primary'])
            
            # Set category labels
            ax.set_xticks(angles[:-1])
            ax.set_xticklabels(categories[:-1], size=12)
            
            # Set y-axis limits
            ax.set_ylim(0, 1)
            ax.set_yticks([0.2, 0.4, 0.6, 0.8, 1.0])
            ax.set_yticklabels(['0.2', '0.4', '0.6', '0.8', '1.0'], size=10)
            
            # Add title
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            plt.title(f"{restaurant_name} Component Scores", size=16, color=self.colors['dark'], pad=20)
            
            # Save figure
            plt.tight_layout()
            plt.savefig(output_path, dpi=300, bbox_inches='tight')
            plt.close()
        except Exception as e:
            logger.error(f"Error generating radar chart: {e}")
    
    def _generate_benchmark_chart(self, grading_results, output_path):
        """
        Generate benchmark comparison bar chart.
        
        Args:
            grading_results (dict): Grading results
            output_path (str): Path to save the chart
        """
        try:
            benchmark_comparisons = grading_results.get('benchmark_comparisons', {})
            
            # Prepare data
            components = []
            scores = []
            benchmarks = []
            
            for component, data in benchmark_comparisons.items():
                components.append(component.replace('_', ' ').title())
                scores.append(data.get('score', 0))
                benchmarks.append(data.get('benchmark_avg', 0))
            
            # Create figure
            fig, ax = plt.subplots(figsize=(12, 8))
            
            # Set background color
            fig.patch.set_facecolor(self.colors['light'])
            ax.set_facecolor(self.colors['light'])
            
            # Set bar width
            bar_width = 0.35
            
            # Set positions
            positions = np.arange(len(components))
            
            # Plot bars
            ax.bar(positions - bar_width/2, scores, bar_width, label='Your Score', color=self.colors['primary'])
            ax.bar(positions + bar_width/2, benchmarks, bar_width, label='Industry Average', color=self.colors['secondary'])
            
            # Add labels and title
            ax.set_xlabel('Component', fontsize=12)
            ax.set_ylabel('Score', fontsize=12)
            
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            ax.set_title(f"{restaurant_name} vs. Industry Benchmarks", fontsize=16, pad=20)
            
            # Set x-axis ticks
            ax.set_xticks(positions)
            ax.set_xticklabels(components, rotation=45, ha='right')
            
            # Set y-axis limits
            ax.set_ylim(0, 1)
            
            # Add legend
            ax.legend()
            
            # Add grid
            ax.grid(axis='y', linestyle='--', alpha=0.7)
            
            # Save figure
            plt.tight_layout()
            plt.savefig(output_path, dpi=300, bbox_inches='tight')
            plt.close()
        except Exception as e:
            logger.error(f"Error generating benchmark chart: {e}")
    
    def _generate_grade_chart(self, grading_results, output_path):
        """
        Generate grade distribution pie chart.
        
        Args:
            grading_results (dict): Grading results
            output_path (str): Path to save the chart
        """
        try:
            component_grades = grading_results.get('component_grades', {})
            
            # Count grade occurrences
            grade_counts = {}
            for grade in component_grades.values():
                if grade not in grade_counts:
                    grade_counts[grade] = 0
                grade_counts[grade] += 1
            
            # Prepare data
            grades = []
            counts = []
            colors = []
            
            for grade, count in sorted(grade_counts.items()):
                grades.append(grade)
                counts.append(count)
                colors.append(self.colors['grades'].get(grade, '#CCCCCC'))
            
            # Create figure
            fig, ax = plt.subplots(figsize=(10, 8))
            
            # Set background color
            fig.patch.set_facecolor(self.colors['light'])
            ax.set_facecolor(self.colors['light'])
            
            # Plot pie chart
            wedges, texts, autotexts = ax.pie(
                counts, 
                labels=grades, 
                autopct='%1.1f%%',
                startangle=90,
                colors=colors,
                wedgeprops={'edgecolor': 'w', 'linewidth': 1}
            )
            
            # Style text
            for text in texts:
                text.set_fontsize(12)
            for autotext in autotexts:
                autotext.set_fontsize(12)
                autotext.set_color('white')
            
            # Add title
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            ax.set_title(f"{restaurant_name} Grade Distribution", fontsize=16, pad=20)
            
            # Equal aspect ratio ensures that pie is drawn as a circle
            ax.axis('equal')
            
            # Save figure
            plt.tight_layout()
            plt.savefig(output_path, dpi=300, bbox_inches='tight')
            plt.close()
        except Exception as e:
            logger.error(f"Error generating grade chart: {e}")
    
    def generate_pdf_report(self, grading_results, output_dir=None):
        """
        Generate PDF report for grading results.
        
        Args:
            grading_results (dict): Grading results
            output_dir (str, optional): Directory to save the report
            
        Returns:
            str: Path to generated PDF report
        """
        try:
            output_dir = output_dir or self.output_dir
            os.makedirs(output_dir, exist_ok=True)
            
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            
            # Generate visualizations
            viz_paths = self.generate_visualizations(grading_results, output_dir)
            
            # Create PDF
            pdf = FPDF()
            pdf.add_page()
            
            # Set font
            pdf.set_font('Arial', 'B', 16)
            
            # Title
            pdf.cell(0, 10, f"Uber Eats Storefront Analysis: {restaurant_name}", 0, 1, 'C')
            pdf.ln(10)
            
            # Overall grade
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, f"Overall Grade: {grading_results.get('overall_grade', 'N/A')}", 0, 1, 'L')
            pdf.set_font('Arial', '', 12)
            pdf.cell(0, 10, f"Score: {grading_results.get('overall_score', 0):.2f} / 1.00", 0, 1, 'L')
            pdf.ln(5)
            
            # Component grades
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, "Component Grades", 0, 1, 'L')
            pdf.ln(5)
            
            pdf.set_font('Arial', '', 12)
            for component, grade in grading_results.get('component_grades', {}).items():
                component_name = component.replace('_', ' ').title()
                score = grading_results.get('component_scores', {}).get(component, 0)
                pdf.cell(0, 10, f"{component_name}: {grade} ({score:.2f} / 1.00)", 0, 1, 'L')
            
            pdf.ln(10)
            
            # Add visualizations
            if viz_paths.get('radar_chart'):
                pdf.add_page()
                pdf.set_font('Arial', 'B', 14)
                pdf.cell(0, 10, "Component Scores", 0, 1, 'C')
                pdf.image(viz_paths.get('radar_chart'), x=10, y=30, w=190)
            
            if viz_paths.get('benchmark_chart'):
                pdf.add_page()
                pdf.set_font('Arial', 'B', 14)
                pdf.cell(0, 10, "Benchmark Comparison", 0, 1, 'C')
                pdf.image(viz_paths.get('benchmark_chart'), x=10, y=30, w=190)
            
            # Key strengths
            pdf.add_page()
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, "Key Strengths", 0, 1, 'L')
            pdf.ln(5)
            
            pdf.set_font('Arial', '', 12)
            for i, strength in enumerate(grading_results.get('strengths', [])[:5]):
                pdf.cell(0, 10, f"{i+1}. {strength.get('text', '')}", 0, 1, 'L')
            
            pdf.ln(10)
            
            # Areas for improvement
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, "Areas for Improvement", 0, 1, 'L')
            pdf.ln(5)
            
            pdf.set_font('Arial', '', 12)
            for i, weakness in enumerate(grading_results.get('weaknesses', [])[:5]):
                pdf.cell(0, 10, f"{i+1}. {weakness.get('text', '')}", 0, 1, 'L')
            
            pdf.ln(10)
            
            # Prioritized recommendations
            pdf.add_page()
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, "Prioritized Recommendations", 0, 1, 'L')
            pdf.ln(5)
            
            pdf.set_font('Arial', '', 12)
            for i, rec in enumerate(grading_results.get('prioritized_recommendations', [])):
                pdf.multi_cell(0, 10, f"{i+1}. {rec.get('recommendation', '')}")
                pdf.ln(5)
            
            # Save PDF
            pdf_path = os.path.join(output_dir, f"{restaurant_name}_Report_{timestamp}.pdf")
            pdf.output(pdf_path)
            
            return pdf_path
        except Exception as e:
            logger.error(f"Error generating PDF report: {e}")
            return None
    
    def generate_presentation(self, grading_results, output_dir=None):
        """
        Generate PowerPoint presentation for grading results.
        
        Args:
            grading_results (dict): Grading results
            output_dir (str, optional): Directory to save the presentation
            
        Returns:
            str: Path to generated presentation
        """
        try:
            if not PPTX_AVAILABLE:
                logger.warning("python-pptx not available, skipping presentation generation")
                return None
                
            logger.info(f"Generating presentation for {grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')}")
            
            output_dir = output_dir or self.output_dir
            os.makedirs(output_dir, exist_ok=True)
            
            restaurant_name = grading_results.get('restaurant_info', {}).get('name', 'Unknown Restaurant')
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            
            # Generate visualizations
            viz_paths = self.generate_visualizations(grading_results, output_dir)
            
            # Create presentation
            prs = Presentation()
            
            # Title slide
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            title.text = f"Uber Eats Storefront Analysis"
            subtitle.text = restaurant_name
            
            # Summary slide
            summary_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(summary_slide_layout)
            title = slide.shapes.title
            content = slide.placeholders[1]
            
            title.text = "Executive Summary"
            
            overall_grade = grading_results.get('overall_grade', 'N/A')
            overall_score = grading_results.get('overall_score', 0)
            
            reviews_count = grading_results.get('analysis_results', {}).get('customer_sentiment', {}).get('metrics', {}).get('review_count', 0)
            avg_rating = grading_results.get('analysis_results', {}).get('customer_sentiment', {}).get('metrics', {}).get('average_rating', 0)
            
            content.text = f"""
            Overall Grade: {overall_grade} ({overall_score:.2f} / 1.00)
            
            Customer Reviews: {reviews_count} reviews with {avg_rating:.1f}/5.0 average rating
            
            Key Strengths:
            - {grading_results.get('strengths', [{'text': 'N/A'}])[0].get('text', 'N/A') if grading_results.get('strengths') else 'N/A'}
            - {grading_results.get('strengths', [{'text': 'N/A'}, {'text': 'N/A'}])[1].get('text', 'N/A') if len(grading_results.get('strengths', [])) > 1 else 'N/A'}
            
            Top Recommendations:
            - {grading_results.get('prioritized_recommendations', [{'recommendation': 'N/A'}])[0].get('recommendation', 'N/A') if grading_results.get('prioritized_recommendations') else 'N/A'}
            - {grading_results.get('prioritized_recommendations', [{'recommendation': 'N/A'}, {'recommendation': 'N/A'}])[1].get('recommendation', 'N/A') if len(grading_results.get('prioritized_recommendations', [])) > 1 else 'N/A'}
            """
            
            # Save presentation
            pptx_path = os.path.join(output_dir, f"{restaurant_name}_Presentation_{timestamp}.pptx")
            prs.save(pptx_path)
            
            return pptx_path
        except Exception as e:
            logger.error(f"Error generating presentation: {e}")
            return None
